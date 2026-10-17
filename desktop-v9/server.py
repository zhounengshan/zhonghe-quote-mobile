import base64
import hashlib
import hmac
import json
import os
import platform
import secrets
import subprocess
import threading
import time
import webbrowser
import sys
import ctypes
import mimetypes
from io import BytesIO
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

# 程序目录与业务数据目录彻底分离：
# Windows/Mac 可以运行不同的程序，但共同读取百度网盘同步的同一套业务数据。
APP_DIR = os.path.dirname(os.path.abspath(sys.executable)) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
RESOURCE_DIR = getattr(sys, '_MEIPASS', APP_DIR)

def user_config_dir():
    if platform.system() == 'Windows':
        return os.path.join(os.environ.get('APPDATA', os.path.expanduser('~')), 'ZhongheBaojiaERP')
    return os.path.join(os.path.expanduser('~'), 'Library', 'Application Support', 'ZhongheBaojiaERP')

CONFIG_DIR = user_config_dir()
DATA_CONFIG_FILE = os.path.join(CONFIG_DIR, 'data_dir.txt')

def default_data_dir():
    # 默认不放在程序目录，避免升级程序时覆盖业务数据。
    return os.path.join(os.path.expanduser('~'), 'Documents', '中禾报价ERP数据')

def resolve_data_dir():
    env_path = os.environ.get('ZHONGHE_ERP_DATA_DIR', '').strip()
    if env_path:
        return os.path.abspath(os.path.expanduser(env_path))
    try:
        if os.path.exists(DATA_CONFIG_FILE):
            path = open(DATA_CONFIG_FILE, 'r', encoding='utf-8').read().strip()
            if path:
                return os.path.abspath(os.path.expanduser(path))
    except Exception:
        pass
    return default_data_dir()

DATA_DIR = resolve_data_dir()
DATA_FILE = os.path.join(DATA_DIR, 'data.json')
AUTH_FILE = os.path.join(DATA_DIR, 'auth.json')
BACKUP_DIR = os.path.join(DATA_DIR, 'backup')
INDEX_FILE = os.path.join(RESOURCE_DIR, 'index.html')
IMAGES_DIR = os.path.join(DATA_DIR, 'images')
LOCK_FILE = os.path.join(CONFIG_DIR, 'ZhongheBaojiaERP.lock')
REVISION_FILE = os.path.join(DATA_DIR, 'data.revision')
SESSION_COOKIE = 'zh_baojia_session'
SESSION_TIMEOUT = 12 * 60 * 60  # 12小时

DEFAULT_DATA = {"customers": [], "quotes": [], "nextCustomerId": 1, "nextQuoteId": 1}
sessions = {}
sessions_lock = threading.Lock()
file_lock = threading.Lock()


def data_revision():
    try:
        h = hashlib.sha256()
        with open(DATA_FILE, 'rb') as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b''):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ''


def save_data_revision(rev=None):
    rev = rev or data_revision()
    try:
        with open(REVISION_FILE, 'w', encoding='utf-8') as f:
            f.write(rev)
    except Exception:
        pass
    return rev


def ensure_data_file():
    if not os.path.exists(DATA_FILE):
        with open(DATA_FILE, 'w', encoding='utf-8') as f:
            json.dump(DEFAULT_DATA, f, ensure_ascii=False, indent=2)


def machine_fingerprint():
    """生成本机指纹。授权文件复制到另一台电脑后，指纹会不一致。"""
    parts = [platform.system(), platform.machine(), platform.node()]
    try:
        if platform.system() == 'Darwin':
            out = subprocess.check_output(
                ['ioreg', '-rd1', '-c', 'IOPlatformExpertDevice'],
                stderr=subprocess.DEVNULL, text=True, timeout=3
            )
            for line in out.splitlines():
                if 'IOPlatformUUID' in line:
                    parts.append(line.split('=', 1)[-1].strip().strip('"'))
                    break
        elif platform.system() == 'Windows':
            out = subprocess.check_output(
                ['powershell', '-NoProfile', '-Command', '(Get-CimInstance Win32_ComputerSystemProduct).UUID'],
                stderr=subprocess.DEVNULL, text=True, timeout=3
            )
            parts.append(out.strip())
        else:
            for path in ('/etc/machine-id', '/var/lib/dbus/machine-id'):
                if os.path.exists(path):
                    parts.append(open(path, encoding='utf-8').read().strip())
                    break
    except Exception:
        pass
    return hashlib.sha256('|'.join(parts).encode('utf-8')).hexdigest()


def hash_password(password, salt=None):
    if salt is None:
        salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt, 210_000)
    return base64.b64encode(salt).decode(), base64.b64encode(digest).decode()


def verify_password(password, salt_b64, digest_b64):
    try:
        salt = base64.b64decode(salt_b64)
        expected = base64.b64decode(digest_b64)
        actual = hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt, 210_000)
        return hmac.compare_digest(actual, expected)
    except Exception:
        return False


def generate_recovery_code():
    return ''.join(secrets.choice('0123456789') for _ in range(12))

def verify_recovery_code(code, auth):
    return bool(auth.get('recovery_salt') and auth.get('recovery_hash') and verify_password(str(code).strip(), auth.get('recovery_salt',''), auth.get('recovery_hash','')))

def ensure_recovery_code(auth):
    if auth.get('recovery_salt') and auth.get('recovery_hash'):
        return auth, None
    code=generate_recovery_code(); salt,digest=hash_password(code)
    auth['recovery_salt']=salt; auth['recovery_hash']=digest; auth['recovery_created_at']=int(time.time())
    with open(AUTH_FILE,'w',encoding='utf-8') as f: json.dump(auth,f,ensure_ascii=False,indent=2)
    return auth, code


def load_auth():
    if not os.path.exists(AUTH_FILE):
        return None
    try:
        with open(AUTH_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return None


def data_is_nonempty():
    try:
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return bool(data.get('customers') or data.get('quotes'))
    except Exception:
        return False


def json_response(handler, status, payload):
    raw = json.dumps(payload, ensure_ascii=False).encode('utf-8')
    handler.send_response(status)
    handler.send_header('Content-Type', 'application/json; charset=utf-8')
    handler.send_header('Content-Length', str(len(raw)))
    handler.send_header('Cache-Control', 'no-store')
    handler.end_headers()
    handler.wfile.write(raw)


def get_session(handler):
    cookie = handler.headers.get('Cookie', '')
    token = None
    for part in cookie.split(';'):
        name, sep, value = part.strip().partition('=')
        if sep and name == SESSION_COOKIE:
            token = value
            break
    if not token:
        return None
    with sessions_lock:
        session = sessions.get(token)
        if not session:
            return None
        if time.time() - session['last_activity'] > SESSION_TIMEOUT:
            sessions.pop(token, None)
            return None
        session['last_activity'] = time.time()
        return session


def require_auth(handler):
    session = get_session(handler)
    if not session:
        json_response(handler, 401, {'error': 'unauthorized'})
        return None
    return session


def create_session(username):
    token = secrets.token_urlsafe(32)
    with sessions_lock:
        sessions[token] = {'username': username, 'last_activity': time.time()}
    return token


def set_session_cookie(handler, token):
    handler.send_header('Set-Cookie', f'{SESSION_COOKIE}={token}; HttpOnly; SameSite=Strict; Path=/; Max-Age={SESSION_TIMEOUT}')


def backup_data(reason='save'):
    ensure_data_file()
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = time.strftime('%Y%m%d_%H%M%S')
    target = os.path.join(BACKUP_DIR, f'data_{stamp}_{reason}.json')
    with open(DATA_FILE, 'rb') as src, open(target, 'wb') as dst:
        dst.write(src.read())
    # 保留最近30份
    files = sorted(
        [os.path.join(BACKUP_DIR, x) for x in os.listdir(BACKUP_DIR) if x.endswith('.json')],
        key=os.path.getmtime,
        reverse=True
    )
    for old in files[30:]:
        try:
            os.remove(old)
        except OSError:
            pass


def read_json_body(handler):
    length = int(handler.headers.get('Content-Length', '0'))
    if length > 20 * 1024 * 1024:
        raise ValueError('请求数据过大')
    return json.loads(handler.rfile.read(length).decode('utf-8'))


def normalize_business_data(data):
    if not isinstance(data, dict):
        data = {}
    data['customers'] = data.get('customers') if isinstance(data.get('customers'), list) else []
    data['quotes'] = data.get('quotes') if isinstance(data.get('quotes'), list) else []
    data['nextCustomerId'] = int(data.get('nextCustomerId') or 1)
    data['nextQuoteId'] = int(data.get('nextQuoteId') or 1)
    data['settings'] = data.get('settings') if isinstance(data.get('settings'), dict) else {}
    return data


def same_record(a, b):
    return json.dumps(a, ensure_ascii=False, sort_keys=True) == json.dumps(b, ensure_ascii=False, sort_keys=True)


def next_available_id(prefix, used, start):
    n = max(1, int(start or 1))
    while f'{prefix}{n}' in used:
        n += 1
    value = f'{prefix}{n}'
    used.add(value)
    return value, n + 1


def max_next_id(records, prefix, current):
    n = int(current or 1)
    for item in records:
        raw = str(item.get('id', '')) if isinstance(item, dict) else ''
        if raw.startswith(prefix) and raw[len(prefix):].isdigit():
            n = max(n, int(raw[len(prefix):]) + 1)
    return n


def merge_customers(current_items, incoming_items, next_id):
    result = [dict(x) for x in current_items if isinstance(x, dict)]
    by_id = {str(x.get('id')): x for x in result if x.get('id')}
    by_key = {str(x.get('key')): x for x in result if x.get('key')}
    used = set(by_id.keys())
    next_id = max_next_id(result + [x for x in incoming_items if isinstance(x, dict)], 'c', next_id)

    for item in incoming_items:
        if not isinstance(item, dict):
            continue
        rec = dict(item)
        key = str(rec.get('key') or '')
        rec_id = str(rec.get('id') or '')
        target = by_key.get(key) if key else None
        if not target and rec_id:
            target = by_id.get(rec_id)

        if target and (not key or target.get('key') == rec.get('key')):
            for field, value in rec.items():
                if field == 'id':
                    continue
                if field == 'lastQuoteDate':
                    if str(value or '') > str(target.get(field) or ''):
                        target[field] = value
                elif value not in (None, '') or target.get(field) in (None, ''):
                    target[field] = value
            continue

        if not rec_id or rec_id in used:
            rec['id'], next_id = next_available_id('c', used, next_id)
        else:
            used.add(rec_id)
        result.append(rec)
        by_id[str(rec.get('id'))] = rec
        if rec.get('key'):
            by_key[str(rec.get('key'))] = rec
    return result, next_id


def quote_identity(q):
    return (
        str(q.get('number') or ''),
        str(q.get('customerName') or ''),
        str(q.get('customerType') or ''),
        str(q.get('customerIdType') or ''),
        str(q.get('customerIdNumber') or ''),
        str(q.get('date') or ''),
    )


def merge_quotes(current_items, incoming_items, next_id):
    result = [dict(x) for x in current_items if isinstance(x, dict)]
    by_id = {str(x.get('id')): x for x in result if x.get('id')}
    used = set(by_id.keys())
    next_id = max_next_id(result + [x for x in incoming_items if isinstance(x, dict)], 'q', next_id)

    for item in incoming_items:
        if not isinstance(item, dict):
            continue
        rec = dict(item)
        rec_id = str(rec.get('id') or '')
        existing = by_id.get(rec_id) if rec_id else None
        if existing:
            if same_record(existing, rec):
                continue
            same_created = str(existing.get('createdAt') or '') == str(rec.get('createdAt') or '')
            same_identity = quote_identity(existing) == quote_identity(rec)
            if same_created or same_identity:
                existing.clear()
                existing.update(rec)
                continue
            rec['id'], next_id = next_available_id('q', used, next_id)
        elif not rec_id or rec_id in used:
            rec['id'], next_id = next_available_id('q', used, next_id)
        else:
            used.add(rec_id)
        result.append(rec)
        by_id[str(rec.get('id'))] = rec
    return result, next_id


def merge_business_data(current_data, incoming_data):
    current = normalize_business_data(current_data)
    incoming = normalize_business_data(incoming_data)
    merged = dict(current)
    customers, next_customer_id = merge_customers(current['customers'], incoming['customers'], max(current['nextCustomerId'], incoming['nextCustomerId']))
    quotes, next_quote_id = merge_quotes(current['quotes'], incoming['quotes'], max(current['nextQuoteId'], incoming['nextQuoteId']))
    merged['customers'] = customers
    merged['quotes'] = quotes
    merged['nextCustomerId'] = max_next_id(customers, 'c', next_customer_id)
    merged['nextQuoteId'] = max_next_id(quotes, 'q', next_quote_id)
    merged['settings'] = dict(current.get('settings') or {})
    merged['settings'].update(incoming.get('settings') or {})
    return merged



def ensure_runtime_dirs():
    os.makedirs(CONFIG_DIR, exist_ok=True)
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(BACKUP_DIR, exist_ok=True)
    os.makedirs(IMAGES_DIR, exist_ok=True)


def safe_image_name(name):
    name = os.path.basename(str(name or ''))
    if not name or name in ('.','..') or '/' in name or '\\' in name:
        raise ValueError('无效的图片文件名')
    return name


def save_image_data(data_url, original_name='image.png'):
    if not isinstance(data_url,str) or not data_url.startswith('data:image/') or ',' not in data_url:
        raise ValueError('无效的图片数据')
    header, enc = data_url.split(',',1)
    mime = header.split(';',1)[0].split(':',1)[1].lower()
    ext = { 'image/jpeg':'.jpg','image/png':'.png','image/gif':'.gif','image/webp':'.webp','image/bmp':'.bmp' }.get(mime, '.png')
    raw=base64.b64decode(enc, validate=True)
    if len(raw)>12*1024*1024: raise ValueError('单张图片不能超过12MB')
    ensure_runtime_dirs(); name=secrets.token_hex(16)+ext
    with open(os.path.join(IMAGES_DIR,name),'wb') as f:f.write(raw)
    return name


def image_bytes(ref):
    if not ref:return None
    if isinstance(ref,str) and ref.startswith('data:image/'):
        try:return base64.b64decode(ref.split(',',1)[1])
        except:return None
    try:
        name=safe_image_name(ref); path=os.path.join(IMAGES_DIR,name)
        if not os.path.isfile(path):return None
        with open(path,'rb') as f:return f.read()
    except:return None


def merged_packing(item):
    details = item.get('packingDetails')
    base = item.get('packing') or ''
    if not isinstance(details, dict):
        return base
    fields = [('productWeight','产品克重（g）'),('boxSpec','箱规（cm）'),('packingQuantity','装箱数量（PCS）'),('grossWeight','毛重（kg）'),('netWeight','净重（kg）')]
    block = item.get('packingGeneratedBlock')
    if block:
        base = base.replace(block, '', 1)
    labels = [label for key, label in fields if key in details]
    lines = [line for line in base.splitlines() if not any(line.strip().startswith((label+'：', label+':')) for label in labels)]
    lines += [label+'：'+str(details[key]) for key, label in fields if details.get(key) is not None and str(details[key]).strip()]
    return '\n'.join(lines).strip()


def export_excel(quote):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    wb=Workbook(); ws=wb.active; ws.title='报价单'
    st=quote.get('settings') or {}; company=st.get('companyName') or '温州市中禾工艺品有限公司'
    ws.merge_cells('A1:M1'); ws['A1']=company; ws['A1'].font=Font(size=18,bold=True); ws['A1'].alignment=Alignment(horizontal='center')
    ws.merge_cells('A2:M2'); ws['A2']='报价单号：%s    报价日期：%s    有效期：%s天（至 %s）'%(quote.get('number',''),quote.get('date',''),quote.get('validityDays',30),quote.get('validUntil','')); ws['A2'].alignment=Alignment(horizontal='center')
    heads=['序号','品名','设计/参考图','材质','规格','工艺要求','数量(pcs)','单价(元/pcs)','模具/版费(元)','合计(元)','工期','装箱数据','备注']
    for c,h in enumerate(heads,1):
        x=ws.cell(5,c,h); x.font=Font(bold=True); x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); x.fill=PatternFill('solid',fgColor='EAF1F8')
    widths=[7,16,20,12,20,26,11,14,16,15,12,20,20]
    for c,w in enumerate(widths,1):ws.column_dimensions[chr(64+c)].width=w
    side=Side(style='thin',color='B7B7B7'); border=Border(left=side,right=side,top=side,bottom=side); row=6; total=0
    multi = any(item.get('quantityTiers') for item in quote.get('items') or [])
    export_items = []
    groups = []
    for product_index, item in enumerate(quote.get('items') or [], 1):
        tiers = [item] + [{**item, **tier} for tier in item.get('quantityTiers') or []]
        groups.append((6+len(export_items),len(tiers)))
        for n, tier in enumerate(tiers, 1):
            export_items.append({**tier, '_productIndex':product_index, 'images':item.get('images',[]) if n==1 else []})
    for i,item in enumerate(export_items,1):
        qty=float(item.get('quantity') or 0); price=float(item.get('price') or 0); mold=float(item.get('moldFee') or 0); sub=qty*price+mold; total+=sub
        vals=[item['_productIndex'],item.get('name') or '-', '',item.get('material') or '-',item.get('spec') or '-',item.get('process') or '-',qty,price,mold,sub,item.get('leadTime') or '-',merged_packing(item) or '-',item.get('remark') or '']
        for c,v in enumerate(vals,1):
            x=ws.cell(row,c,v); x.border=border; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ims=[]
        for obj in (item.get('images') or [])[:6]:
            ref=obj.get('filename') if isinstance(obj,dict) else obj; raw=image_bytes(ref)
            if raw:
                try:
                    im=XLImage(BytesIO(raw)); im.width=min(120,im.width); im.height=min(90,im.height); ims.append(im)
                except:pass
        for j,im in enumerate(ims): im.anchor='C'+str(row); ws.add_image(im); break
        ws.row_dimensions[row].height=85; row+=1
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9); ws.cell(row,1,'各档位为备选方案，金额不累加' if multi else '报价总额').font=Font(bold=True,size=14); ws.cell(row,1).alignment=Alignment(horizontal='right'); ws.cell(row,10,None if multi or quote.get('showTotal') is False else total).font=Font(bold=True,size=14); ws.cell(row,10).number_format='¥#,##0.00'; ws.cell(row,10).alignment=Alignment(horizontal='center');
    for c in range(1,14):ws.cell(row,c).border=border
    for start, count in groups:
        if count>1:
            for col in [1,2,3,4,5,6,12,13]:ws.merge_cells(start_row=start,start_column=col,end_row=start+count-1,end_column=col)
    ws.freeze_panes='A6'; return wb


def acquire_single_instance():
    if os.name != 'nt': return None
    handle=ctypes.windll.kernel32.CreateMutexW(None, False, 'Global\\ZhongheBaojiaERP_SingleInstance')
    if not handle: return None
    ERROR_ALREADY_EXISTS=183
    if ctypes.windll.kernel32.GetLastError()==ERROR_ALREADY_EXISTS:
        ctypes.windll.kernel32.CloseHandle(handle); return False
    return handle


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        path = urlparse(self.path).path

        if path in ('/efficiency.css', '/efficiency.js', '/business.css', '/business.js'):
            filename = os.path.join(RESOURCE_DIR, path[1:])
            with open(filename, 'rb') as f:
                raw = f.read()
            self.send_response(200)
            self.send_header('Content-Type', 'text/css; charset=utf-8' if path.endswith('.css') else 'application/javascript; charset=utf-8')
            self.send_header('Content-Length', str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)
            return

        if path == '/':
            try:
                with open(INDEX_FILE, 'rb') as f:
                    raw = f.read()
                self.send_response(200)
                self.send_header('Content-Type', 'text/html; charset=utf-8')
                self.send_header('Content-Length', str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
            except Exception as e:
                json_response(self, 500, {'error': str(e)})
            return

        if path == '/api/auth/status':
            auth = load_auth()
            session = get_session(self)
            # 兼容旧版设备指纹：不再因为指纹变化阻断登录。
            json_response(self, 200, {
                'setupRequired': auth is None,
                'authorizedDevice': True,
                'loggedIn': bool(session),
                'username': session.get('username') if session else None
            })
            return

        if path == '/api/image':
            if require_auth(self) is None: return
            try:
                name=safe_image_name((parse_qs(urlparse(self.path).query).get('name') or [''])[0]); path2=os.path.join(IMAGES_DIR,name)
                if not os.path.isfile(path2): self.send_response(404); self.end_headers(); return
                with open(path2,'rb') as f: raw=f.read()
                self.send_response(200); self.send_header('Content-Type',mimetypes.guess_type(path2)[0] or 'application/octet-stream'); self.send_header('Cache-Control','no-store'); self.send_header('Content-Length',str(len(raw))); self.end_headers(); self.wfile.write(raw)
            except Exception as e: json_response(self,400,{'error':str(e)})
            return

        if path == '/api/export/excel':
            if require_auth(self) is None: return
            try:
                qid=(parse_qs(urlparse(self.path).query).get('id') or [''])[0]; ensure_data_file()
                with open(DATA_FILE,'r',encoding='utf-8') as f:data=json.load(f)
                q=next((x for x in data.get('quotes',[]) if str(x.get('id'))==str(qid)),None)
                if not q: json_response(self,404,{'error':'报价单不存在'}); return
                q['settings']=data.get('settings') or {}; wb=export_excel(q); out=BytesIO(); wb.save(out); raw=out.getvalue()
                self.send_response(200); self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); self.send_header('Content-Disposition','attachment; filename="quote.xlsx"'); self.send_header('Content-Length',str(len(raw))); self.end_headers(); self.wfile.write(raw)
            except Exception as e: json_response(self,500,{'error':'Excel导出失败：'+str(e)})
            return

        if path == '/api/data':
            if require_auth(self) is None:
                return
            try:
                ensure_data_file()
                with file_lock, open(DATA_FILE, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                raw = json.dumps(data, ensure_ascii=False).encode('utf-8')
                rev = data_revision()
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(raw)))
                self.send_header('Cache-Control', 'no-store')
                self.send_header('X-Data-Revision', rev)
                self.end_headers()
                self.wfile.write(raw)
            except Exception as e:
                json_response(self, 500, {'error': str(e)})
            return

        if path == '/api/backups':
            if require_auth(self) is None:
                return
            os.makedirs(BACKUP_DIR, exist_ok=True)
            files = sorted([x for x in os.listdir(BACKUP_DIR) if x.endswith('.json')], reverse=True)
            json_response(self, 200, {'files': files[:30]})
            return

        self.send_response(404)
        self.end_headers()

    def do_POST(self):
        path = urlparse(self.path).path

        if path == '/api/auth/setup':
            if load_auth() is not None:
                json_response(self, 409, {'error': 'already_setup'})
                return
            # 如果已有业务数据但授权文件丢失，不允许通过重新初始化绕过授权。
            if data_is_nonempty():
                json_response(self, 403, {'error': 'authorization_missing_for_existing_data'})
                return
            try:
                body = read_json_body(self)
                username = str(body.get('username', '')).strip()
                password = str(body.get('password', ''))
                if len(username) < 3:
                    raise ValueError('用户名至少3位')
                if len(password) < 8:
                    raise ValueError('密码至少8位')
                salt, digest = hash_password(password)
                recovery_code = generate_recovery_code()
                recovery_salt, recovery_hash = hash_password(recovery_code)
                auth = {
                    'username': username,
                    'salt': salt,
                    'password_hash': digest,
                    'device_id': machine_fingerprint(),
                    'created_at': int(time.time()),
                    'recovery_salt': recovery_salt,
                    'recovery_hash': recovery_hash,
                    'recovery_created_at': int(time.time())
                }
                with open(AUTH_FILE, 'w', encoding='utf-8') as f:
                    json.dump(auth, f, ensure_ascii=False, indent=2)
                auth['device_id'] = machine_fingerprint()
                try:
                    with open(AUTH_FILE, 'w', encoding='utf-8') as f:
                        json.dump(auth, f, ensure_ascii=False, indent=2)
                except Exception:
                    pass
                token = create_session(username)
                raw = json.dumps({'status': 'ok', 'username': username, 'recoveryCode': recovery_code}, ensure_ascii=False).encode('utf-8')
                self.send_response(200)
                set_session_cookie(self, token)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
            except Exception as e:
                json_response(self, 400, {'error': str(e)})
            return

        if path == '/api/auth/login':
            try:
                auth = load_auth()
                if not auth:
                    json_response(self, 409, {'error': 'setup_required'})
                    return
                body = read_json_body(self)
                username = str(body.get('username', '')).strip()
                password = str(body.get('password', ''))
                if username != auth.get('username') or not verify_password(password, auth.get('salt', ''), auth.get('password_hash', '')):
                    time.sleep(0.35)
                    json_response(self, 401, {'error': 'invalid_credentials'})
                    return
                token = create_session(username)
                raw = json.dumps({'status': 'ok', 'username': username}, ensure_ascii=False).encode('utf-8')
                self.send_response(200)
                set_session_cookie(self, token)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
            except Exception as e:
                json_response(self, 400, {'error': str(e)})
            return

        if path == '/api/auth/recovery-code':
            auth=load_auth()
            if not auth: json_response(self,409,{'error':'setup_required'}); return
            if require_auth(self) is None: return
            try:
                auth,new_code=ensure_recovery_code(auth)
                json_response(self,200,{'recoveryCode':new_code or '已生成（请以首次生成时保存的找回码为准）'})
            except Exception as e: json_response(self,500,{'error':str(e)})
            return

        if path == '/api/auth/recover-password':
            try:
                auth=load_auth()
                if not auth: json_response(self,409,{'error':'setup_required'}); return
                body=read_json_body(self); username=str(body.get('username','')).strip(); code=str(body.get('recoveryCode','')).strip(); newp=str(body.get('newPassword',''))
                if username != auth.get('username') or not verify_recovery_code(code,auth):
                    time.sleep(0.35); json_response(self,401,{'error':'找回码错误或用户名错误'}); return
                if len(newp)<8: json_response(self,400,{'error':'新密码至少8位'}); return
                salt,digest=hash_password(newp); auth['salt']=salt; auth['password_hash']=digest; auth['password_changed_at']=int(time.time())
                with open(AUTH_FILE,'w',encoding='utf-8') as f: json.dump(auth,f,ensure_ascii=False,indent=2)
                json_response(self,200,{'status':'ok'})
            except Exception as e: json_response(self,400,{'error':str(e)})
            return

        if path == '/api/auth/change-password':
            if require_auth(self) is None: return
            try:
                auth=load_auth(); body=read_json_body(self); oldp=str(body.get('oldPassword','')); newp=str(body.get('newPassword',''))
                if len(newp)<8: json_response(self,400,{'error':'新密码至少8位'}); return
                if not auth or not verify_password(oldp,auth.get('salt',''),auth.get('password_hash','')):
                    json_response(self,401,{'error':'原密码错误'}); return
                salt,digest=hash_password(newp); auth['salt']=salt; auth['password_hash']=digest; auth['password_changed_at']=int(time.time())
                with open(AUTH_FILE,'w',encoding='utf-8') as f: json.dump(auth,f,ensure_ascii=False,indent=2)
                json_response(self,200,{'status':'ok'})
            except Exception as e: json_response(self,400,{'error':str(e)})
            return

        if path == '/api/auth/logout':
            cookie = self.headers.get('Cookie', '')
            token = None
            for part in cookie.split(';'):
                name, sep, value = part.strip().partition('=')
                if sep and name == SESSION_COOKIE:
                    token = value
                    break
            if token:
                with sessions_lock:
                    sessions.pop(token, None)
            raw = b'{"status":"ok"}'
            self.send_response(200)
            self.send_header('Set-Cookie', f'{SESSION_COOKIE}=; HttpOnly; SameSite=Strict; Path=/; Max-Age=0')
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)
            return

        if path == '/api/image':
            if require_auth(self) is None: return
            try:
                body=read_json_body(self); name=save_image_data(body.get('data',''),body.get('filename','image.png')); json_response(self,200,{'status':'ok','name':name})
            except Exception as e: json_response(self,400,{'error':str(e)})
            return

        if path == '/api/data':
            if require_auth(self) is None:
                return
            try:
                data = read_json_body(self)
                if not isinstance(data, dict) or not isinstance(data.get('customers'), list) or not isinstance(data.get('quotes'), list):
                    raise ValueError('数据格式无效')
                expected_revision = self.headers.get('X-Data-Revision', '').strip()
                with file_lock:
                    current_revision = data_revision()
                    saved_by_merge = False
                    if expected_revision and current_revision and expected_revision != current_revision:
                        ensure_data_file()
                        with open(DATA_FILE, 'r', encoding='utf-8') as f:
                            current_data = json.load(f)
                        data = merge_business_data(current_data, data)
                        saved_by_merge = True
                    if os.path.exists(DATA_FILE):
                        backup_data('before_merge_save' if saved_by_merge else 'before_save')
                    tmp = DATA_FILE + '.tmp'
                    with open(tmp, 'w', encoding='utf-8') as f:
                        json.dump(data, f, ensure_ascii=False, indent=2)
                        f.flush()
                        os.fsync(f.fileno())
                    os.replace(tmp, DATA_FILE)
                    with open(DATA_FILE, 'r', encoding='utf-8') as f:
                        saved_data = json.load(f)
                    rev = save_data_revision(data_revision())
                raw = json.dumps({'status': 'ok', 'data': saved_data, 'revision': rev, 'merged': saved_by_merge}, ensure_ascii=False).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(raw)))
                self.send_header('Cache-Control', 'no-store')
                self.send_header('X-Data-Revision', rev)
                self.end_headers()
                self.wfile.write(raw)
            except Exception as e:
                json_response(self, 500, {'error': str(e)})
            return

        if path == '/api/backup':
            if require_auth(self) is None:
                return
            try:
                backup_data('manual')
                json_response(self, 200, {'status': 'ok'})
            except Exception as e:
                json_response(self, 500, {'error': str(e)})
            return

        self.send_response(404)
        self.end_headers()

    def log_message(self, format, *args):
        pass


def run_desktop_window(server, port):
    """优先使用系统 WebView，确保 APP 在独立窗口中运行而不是跳转浏览器。"""
    try:
        import webview
    except ImportError:
        return False

    threading.Thread(target=server.serve_forever, daemon=True).start()
    window = webview.create_window(
        '中禾报价 ERP V9',
        f'http://127.0.0.1:{port}',
        width=1440,
        height=900,
        min_size=(1024, 700),
        background_color='#eef0f2'
    )
    try:
        webview.start(debug=False, private_mode=False, storage_path=CONFIG_DIR)
    finally:
        server.shutdown()
    return window is not None


if __name__ == '__main__':
    ensure_runtime_dirs(); ensure_data_file()
    instance=acquire_single_instance()
    if instance is False:
        webbrowser.open('http://127.0.0.1:8000')
        sys.exit(0)
    port=8000
    try:
        server=HTTPServer(('127.0.0.1',port),Handler)
    except OSError:
        sys.exit(1)
    try:
        if not run_desktop_window(server, port):
            import pystray
            from PIL import Image
            ico=os.path.join(RESOURCE_DIR,'app.ico')
            tray_img=Image.open(ico).convert('RGBA') if os.path.exists(ico) else Image.new('RGBA',(64,64),'white')
            def open_app(icon,item): webbrowser.open(f'http://127.0.0.1:{port}')
            def quit_app(icon,item):
                icon.stop(); threading.Thread(target=server.shutdown,daemon=True).start()
            tray=pystray.Icon('ZhongheBaojiaERP',tray_img,'中禾报价 ERP',pystray.Menu(pystray.MenuItem('打开报价系统',open_app,default=True),pystray.MenuItem('退出报价系统',quit_app)))
            threading.Thread(target=server.serve_forever,daemon=True).start(); threading.Timer(0.8,lambda:webbrowser.open(f'http://127.0.0.1:{port}')).start(); tray.run()
    except Exception as exc:
        print(f'桌面窗口启动失败：{exc}', file=sys.stderr)
        threading.Timer(0.8,lambda:webbrowser.open(f'http://127.0.0.1:{port}')).start(); server.serve_forever()
    finally:
        try: server.server_close()
        except: pass
        if instance and instance is not False:
            try: ctypes.windll.kernel32.CloseHandle(instance)
            except: pass
